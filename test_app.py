import unittest, os, io, uuid, base64, importlib.util, time
from unittest.mock import Mock, patch
from datetime import date
from openpyxl import load_workbook
from engine import calculate, ValidationError, day, add_month
from reports import pdf_data, xlsx_data, csv_data
import mailer

def sample(**kw):
    c=dict(kind='personal',principal=100000,rate=12,months=60,start='2026-10-27',first='2026-11-27',extra=0,service=69,initiation=0,vat=.15,balloon=0,insurance='none',insurance_rate=0,premium=0,rate_type='variable',irregular=[],fees=[])
    return {**c,**kw}

def contact(**kw):return {**dict(name='Test Client',email='client@example.com',mobile='',followup=False,consent=True),**kw}

class EngineTests(unittest.TestCase):
    def test_zero_interest(self):self.assertAlmostEqual(calculate(sample(rate=0,principal=12000,months=12,service=0))['total_paid'],12000)
    def test_pmt(self):self.assertAlmostEqual(calculate(sample())['instalment'],2224.44476849,5)
    def test_month_end(self):self.assertEqual(add_month(day('2028-01-31'),1),date(2028,2,29));self.assertEqual(add_month(day('2028-01-31'),2),date(2028,3,31))
    def test_bad_date(self):
        with self.assertRaises(ValidationError):calculate(sample(first='2026-02-30'))
    def test_whole_term(self):
        with self.assertRaises(ValidationError):calculate(sample(months=6.5))
    def test_negative_nan(self):
        for kw in ({'extra':-1},{'rate':float('nan')},{'principal':0}):
            with self.assertRaises(ValidationError):calculate(sample(**kw))
    def test_midmonth_payment(self):
        x=next(x for x in calculate(sample(irregular=[dict(date='2027-01-15',amount=10000)]))['rows'] if x['date']=='2027-01-15');self.assertAlmostEqual(x['payment'],10000);self.assertGreater(x['capital_paid'],0)
    def test_same_day_interest_once(self):
        x=calculate(sample(irregular=[dict(date='2026-11-27',amount=1000)]))['rows'][0];self.assertAlmostEqual(x['interest'],100000*.12*31/365);self.assertAlmostEqual(x['payment'],calculate(sample())['instalment']+69+1000)
    def test_extra_saves(self):self.assertGreater(calculate(sample(extra=500))['interest_saved'],0)
    def test_initiation_upfront(self):
        r=calculate(sample(initiation=1207.5));self.assertAlmostEqual(r['total_paid'],100000+r['total_interest']+r['total_fees'])
    def test_initiation_financed(self):
        r=calculate(sample(initiation=1207.5,finance=True));self.assertAlmostEqual(r['total_paid'],100000+r['total_interest']+r['total_fees'])
    def test_service_cap(self):
        calculate(sample(service=69))
        with self.assertRaises(ValidationError):calculate(sample(service=69.01))
    def test_initiation_cap(self):
        with self.assertRaises(ValidationError):calculate(sample(initiation=1208))
    def test_insurance(self):
        self.assertGreater(calculate(sample(insurance='reducing',insurance_rate=4.5))['total_insurance'],0)
        with self.assertRaises(ValidationError):calculate(sample(insurance='reducing',insurance_rate=4.51))
    def test_home_insurance(self):
        with self.assertRaises(ValidationError):calculate(sample(kind='home',insurance='reducing',insurance_rate=2.01))
    def test_repo(self):
        with self.assertRaises(ValidationError):calculate(sample(repo=7,rate=29))
    def test_before_first_quote(self):self.assertAlmostEqual(calculate(sample(quote='2026-11-01'))['quote']['interest'],100000*.12*5/365)
    def test_before_start_quote(self):
        with self.assertRaises(ValidationError):calculate(sample(quote='2026-10-26'))
    def test_variable_large_charge(self):self.assertGreater(calculate(sample(kind='home',quote='2027-01-15'))['quote']['termination'],0)
    def test_notice(self):self.assertEqual(calculate(sample(kind='home',quote='2027-02-01',notice='2026-11-01'))['quote']['termination'],0)
    def test_fixed_unknown(self):self.assertIsNone(calculate(sample(kind='home',rate_type='fixed',quote='2027-01-15'))['quote']['total'])
    def test_fixed_known(self):
        q=calculate(sample(kind='home',rate_type='fixed',quote='2027-01-15',fixed_charge=500))['quote'];self.assertAlmostEqual(q['total']-q['subtotal'],500)
    def test_original_large(self):self.assertGreater(calculate(sample(original_large=True,quote='2027-01-15'))['quote']['termination'],0)
    def test_quote_paid_off(self):self.assertEqual(calculate(sample(rate=0,months=1,quote='2027-02-01'))['quote']['total'],0)
    def test_balloon(self):self.assertGreater(calculate(sample(kind='vehicle',balloon=30000))['final_payment'],30000)
    def test_fee_review(self):
        with self.assertRaises(ValidationError):calculate(sample(fees=[dict(date='2027-01-15',kind='collection',amount=100)]))
    def test_fee_accounting(self):
        r=calculate(sample(fees=[dict(date='2027-01-15',kind='collection',amount=100,verified=True)]));self.assertAlmostEqual(r['total_paid'],100000+r['total_interest']+r['total_fees'])
    def test_150_generated_cashflows(self):
        for i in range(1,151):
            c=sample(principal=1000+i*1000,rate=i%30,months=6+i%120,extra=i*3,service=0);r=calculate(c);self.assertAlmostEqual(r['total_paid'],c['principal']+r['total_interest']+r['total_fees'],4)

class ReportTests(unittest.TestCase):
    def setUp(self):self.c=sample();self.r=calculate(self.c);self.p=contact()
    def test_pdf(self):
        a=pdf_data(self.c,self.r,self.p,'reference');b=pdf_data(self.c,self.r,self.p,'reference');self.assertTrue(a.startswith(b'%PDF'));self.assertEqual(a,b)
    def test_xlsx(self):
        a=xlsx_data(self.c,self.r,self.p,'reference');time.sleep(1.05);b=xlsx_data(self.c,self.r,self.p,'reference');self.assertEqual(a,b);w=load_workbook(io.BytesIO(a));self.assertIn('Amortisation',w.sheetnames);self.assertGreater(w['Amortisation'].max_row,1)
    def test_no_formula_injection(self):
        p=contact(name='=1+1');w=load_workbook(io.BytesIO(xlsx_data(self.c,self.r,p,'r')));self.assertEqual(w['Summary']['B2'].data_type,'s')
    def test_csv(self):self.assertIn(b'capital_paid',csv_data(self.r))

class MailTests(unittest.TestCase):
    def setUp(self):
        self.env=patch.dict(os.environ,{'FROM_EMAIL':'SPWM <reports@example.com>','RESEND_API_KEY':'test-key','PUBLIC_URL':'https://calc.example.com','TURNSTILE_SECRET_KEY':'test-secret'});self.env.start();mailer.COUNTS.clear()
    def tearDown(self):self.env.stop()
    def test_consent(self):
        with self.assertRaises(ValidationError):mailer.validate_person(contact(consent=False))
    def test_email_validation(self):
        with self.assertRaises(ValidationError):mailer.validate_person(contact(email='bad\n@example.com'))
    @patch.object(mailer.requests,'post')
    def test_pdf_to_client_excel_only_to_spwm(self,post):
        post.return_value=Mock();post.return_value.json.return_value={'id':'test'};mailer.send_reports(sample(),contact(),str(uuid.uuid4()));a,b=[x.kwargs['json'] for x in post.call_args_list];self.assertEqual(a['to'],['client@example.com']);self.assertEqual(len(a['attachments']),1);self.assertEqual(a['attachments'][0]['filename'],'SPWM-summary.pdf');self.assertEqual(b['to'],['webmaster@spwm.co.za']);self.assertEqual(len(b['attachments']),2);self.assertTrue(base64.b64decode(a['attachments'][0]['content']).startswith(b'%PDF'))
    @patch.object(mailer.requests,'post')
    def test_stable_retry_keys(self,post):
        post.return_value=Mock();post.return_value.json.return_value={'id':'x'};ref=str(uuid.uuid4());mailer.send_reports(sample(),contact(),ref);first=post.call_args_list[:];mailer.send_reports(sample(),contact(),ref)
        for a,b in zip(first,post.call_args_list[2:]):self.assertEqual(a.kwargs,b.kwargs)
    @patch.object(mailer.requests,'post')
    def test_partial_failure(self,post):
        ok=Mock();ok.json.return_value={'id':'x'};post.side_effect=[ok,mailer.requests.Timeout()]
        with self.assertRaises(mailer.DeliveryError) as ctx:mailer.send_reports(sample(),contact(),str(uuid.uuid4()))
        self.assertTrue(ctx.exception.client_accepted)
    def test_throttle(self):
        self.assertTrue(mailer.allowed('a@example.com'));self.assertTrue(mailer.allowed('a@example.com'));self.assertTrue(mailer.allowed('a@example.com'));self.assertFalse(mailer.allowed('a@example.com'))
    @patch.object(mailer.requests,'post')
    def test_bot_hostname_action(self,post):
        post.return_value.json.return_value={'success':True,'hostname':'calc.example.com','action':'report'};self.assertTrue(mailer.verify_bot('test'));post.return_value.json.return_value={'success':True,'hostname':'evil.example.com','action':'report'};self.assertFalse(mailer.verify_bot('test'))

HAS_FLASK=importlib.util.find_spec('flask') is not None
@unittest.skipUnless(HAS_FLASK,'Flask is not installed in this test environment.')
class WebTests(unittest.TestCase):
    def setUp(self):
        import app
        self.web=app;self.env=patch.dict(os.environ,{'PUBLIC_URL':'https://calc.example.com','ENABLE_EMAIL':'false'});self.env.start();self.client=app.app.test_client();self.origin={'Origin':'https://calc.example.com'}
    def tearDown(self):self.env.stop()
    def test_home(self):self.assertEqual(self.client.get('/').status_code,200)
    def test_health(self):self.assertEqual(self.client.get('/health').json['version'],'3.0.0')
    def test_calculate(self):self.assertEqual(self.client.post('/calculate',json={'loan':sample()},headers=self.origin).status_code,200)
    def test_reject_origin(self):self.assertEqual(self.client.post('/calculate',json={'loan':sample()}).status_code,403)
    def test_csv(self):self.assertIn(b'capital_paid',self.client.post('/csv',json={'loan':sample()},headers=self.origin).data)
    def test_email_disabled(self):self.assertEqual(self.client.post('/email',json={},headers=self.origin).status_code,503)
    def test_no_pdf_download(self):self.assertEqual(self.client.get('/pdf').status_code,404)
    def test_email_route(self):
        with patch.object(mailer,'configured',return_value=True),patch.object(mailer,'verify_bot',return_value=True),patch.object(mailer,'allowed',return_value=True),patch.object(mailer,'send_reports',return_value={'reference':'test'}):
            r=self.client.post('/email',json={'loan':sample(),'person':contact(),'token':'t','request_id':str(uuid.uuid4())},headers=self.origin);self.assertEqual(r.status_code,200)

if __name__=='__main__':unittest.main(verbosity=2)

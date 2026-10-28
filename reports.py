"""In-memory PDF, Excel and CSV reports. No public PDF-download endpoint."""
import io, csv, json, zipfile, re
from pathlib import Path
from datetime import datetime
from xml.sax.saxutils import escape
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from openpyxl import Workbook

METRICS=[('Monthly instalment excluding fees','instalment'),('Final payment including residual/balloon','final_payment'),('Total paid','total_paid'),('Total interest','total_interest'),('Fees including insurance','total_fees'),('Insurance already included in fees','total_insurance'),('Interest saved versus no extras','interest_saved')]

def csv_data(r):
    out=io.StringIO();keys=list(r['rows'][0]) if r['rows'] else ['date'];w=csv.DictWriter(out,keys);w.writeheader()
    for row in r['rows']:w.writerow({k:round(v,2) if isinstance(v,float) else v for k,v in row.items()})
    return out.getvalue().encode('utf-8-sig')

def pdf_data(c,r,person,reference):
    out=io.BytesIO();styles=getSampleStyleSheet();story=[]
    def p(text):return Paragraph(escape(str(text)),styles['Normal'])
    logo=Path(__file__).with_name('logo.jpg')
    if logo.is_file():
        image=Image(str(logo));image.drawHeight=image.imageHeight*160/image.imageWidth;image.drawWidth=160;image.hAlign='LEFT';story.extend([image,Spacer(1,12)])
    story += [Paragraph('SPWM Loan Payoff Summary',styles['Title']),p('Reference: '+reference),p('Prepared for: '+person['name']),p('Email: '+person['email']),p('Mobile: '+person.get('mobile','')),Spacer(1,12)]
    pairs=[('Loan type',c['kind']),('Principal at balance date',f"R{float(c['principal']):,.2f}"),('Balance date',c['start']),('First payment',c['first']),('Annual rate',str(c['rate'])+'%'),('Remaining months',str(c['months']))]
    pairs += [(label,f'R{r[key]:,.2f}') for label,key in METRICS];pairs.append(('Payoff',r['payoff'] or 'Not cleared'))
    if r['quote']:
        for key,v in r['quote'].items():pairs.append(('Settlement '+key,'Lender-confirmed charge required' if v is None else f'R{v:,.2f}' if isinstance(v,(float,int)) else v))
    table=Table([[p(a),p(b)] for a,b in pairs],colWidths=[245,265]);table.setStyle(TableStyle([('LINEABOVE',(0,0),(-1,0),2,colors.HexColor('#ce2029')),('BACKGROUND',(0,0),(0,-1),colors.whitesmoke),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),6)]));story.extend([table,Spacer(1,12)])
    for w in r['warnings']:story.extend([p(w),Spacer(1,5)])
    story.extend([p('Actual/365 simple daily interest; entered rate held constant. Extras settle interest, fees, principal in that order. Final payment clears residual and balloon. Monthly costs billed on scheduled dates; no first-period proration. Settlement is after same-day payments.'),Spacer(1,8),p('Estimate only: not financial advice, credit approval, legal certification or a lender settlement quotation. Reference rules need current review. No section 103(5) default reconstruction, court/collection tariffs, debt review, conveyancing or bond-cancellation attorney fees.'),Spacer(1,12),Paragraph('Full amortisation schedule',styles['Heading2'])])
    data=[['Date','Event','Interest','Fees','Paid','Principal','Total owed']]
    for x in r['rows']:data.append([x['date'],p(x['event'])]+[f'{x[k]:,.2f}' for k in ('interest','fees','payment','principal','balance')])
    table=Table(data,colWidths=[55,110,60,50,70,80,80],repeatRows=1);table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.black),('TEXTCOLOR',(0,0),(-1,0),colors.white),('LINEBELOW',(0,0),(-1,0),2,colors.HexColor('#ce2029')),('FONTSIZE',(0,0),(-1,-1),7),('VALIGN',(0,0),(-1,-1),'TOP'),('ALIGN',(2,1),(-1,-1),'RIGHT')]));story.extend([table,Spacer(1,12),Paragraph('Calculation inputs and dated payments / fees',styles['Heading2'])])
    for key,value in sorted(c.items()):story.append(p(key+': '+json.dumps(value,ensure_ascii=True)))
    story.extend([Spacer(1,10),p('Specific follow-up requested: '+str(person.get('followup') is True)),p('Requested PDF sent to client and SPWM. Excel version sent only to webmaster@spwm.co.za.'),p('SPWM 3.0.0 | webmaster@spwm.co.za')])
    SimpleDocTemplate(out,pagesize=A4,leftMargin=25,rightMargin=25,topMargin=28,bottomMargin=28,invariant=1).build(story)
    return out.getvalue()

def xlsx_data(c,r,person,reference):
    wb=Workbook();ws=wb.active;ws.title='Summary';ws.append(['SPWM calculation',reference]);ws.append(['Name',person['name']]);ws.append(['Email',person['email']]);ws.append(['Payoff',r['payoff']])
    for label,key in METRICS:ws.append([label,r[key]])
    if r['quote']:
        for k,v in r['quote'].items():ws.append(['Settlement '+k,v])
    schedule=wb.create_sheet('Amortisation');keys=list(r['rows'][0]) if r['rows'] else ['date'];schedule.append(keys)
    for row in r['rows']:schedule.append([round(row[k],2) if isinstance(row[k],float) else row[k] for k in keys])
    inputs=wb.create_sheet('Inputs');inputs.append(['Field','Value'])
    for k,v in sorted(c.items()):inputs.append([k,json.dumps(v,ensure_ascii=True)])
    notes=wb.create_sheet('Assumptions');notes.append(['Estimate only; not lender quote. Internal SPWM copy only.'])
    for w in r['warnings']:notes.append([w])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.column_dimensions['A'].width=45;sheet.column_dimensions['B'].width=45
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
    wb.properties.created=wb.properties.modified=datetime(2000,1,1)
    raw=io.BytesIO();wb.save(raw)
    # Canonical ZIP metadata makes unchanged retry attachments deterministic.
    out=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as src,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            data=src.read(name)
            if name=='docProps/core.xml':data=re.sub(rb'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)',rb'\g<1>2000-01-01T00:00:00Z\g<2>',data)
            info=zipfile.ZipInfo(name,(2000,1,1,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;dst.writestr(info,data)
    return out.getvalue()
